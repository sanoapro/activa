#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Utilidades de lectura y respaldo de la Base Maestra.

Existe para no repetir las trampas conocidas de estos .xlsx:
  - `False` es una celda llena, no vacía
  - el encabezado no siempre está en la fila 1
  - hay filas de una sola celda entre los datos
  - el nombre del archivo lleva acento y varía

Uso:
    python base.py info
    python base.py cols [hoja]
    python base.py respaldar <fase>
    python base.py contar <hoja> [columna=valor]
"""
import glob
import os
import shutil
import sys
import unicodedata
from datetime import date

PATRON = "0. Base Maestra*.xlsx"
# Sufijos de respaldo: no son la base viva.
RESPALDOS = (".pre-", ".bak")
COLUMNAS_ANCLA = ("ID", "Colegio")


def _raiz():
    """Directorio que contiene la base, buscando hacia arriba desde el script.

    No se cuenta un número fijo de niveles: la skill puede quedar instalada en el
    repo del sitio (`Presentacion-activa/.claude/`) o junto a los datos, y el .xlsx
    vive en `Base de datos - activa/`. Se sube hasta encontrarlo.
    """
    d = os.path.abspath(os.path.dirname(__file__))
    while True:
        if glob.glob(os.path.join(d, PATRON)):
            return d
        padre = os.path.dirname(d)
        if padre == d:
            raise SystemExit(
                "No encontré ningún '0. Base Maestra*.xlsx' subiendo desde "
                + os.path.dirname(os.path.abspath(__file__)))
        d = padre


RAIZ = None  # se resuelve perezosamente en ruta_base()


def ruta_base():
    """Ruta del .xlsx vivo, ignorando respaldos. Falla ruidosamente si hay ambigüedad."""
    global RAIZ
    if RAIZ is None:
        RAIZ = _raiz()
    todos = glob.glob(os.path.join(RAIZ, PATRON))
    vivos = [f for f in todos
             if not any(s in os.path.basename(f) for s in RESPALDOS)]
    if not vivos:
        raise SystemExit(f"No encontré la base viva en {RAIZ}\n"
                         f"Candidatos: {[os.path.basename(f) for f in todos]}")
    if len(vivos) > 1:
        raise SystemExit("Hay más de una base viva, no adivino cuál:\n  " +
                         "\n  ".join(os.path.basename(f) for f in vivos))
    return vivos[0]


def _norm(s):
    if s is None:
        return ""
    s = unicodedata.normalize("NFKD", str(s))
    return "".join(c for c in s if not unicodedata.combining(c)).strip().lower()


def fila_encabezado(ws, limite=15):
    """Localiza la fila del encabezado buscando una columna ancla.

    No asume la fila 1: en los archivos de origen el encabezado se mueve.
    """
    anclas = {_norm(c) for c in COLUMNAS_ANCLA}
    for i, fila in enumerate(ws.iter_rows(min_row=1, max_row=limite,
                                          values_only=True), start=1):
        if anclas & {_norm(c) for c in fila if c is not None}:
            return i
    return 1


def vacia(v):
    """True solo si la celda está realmente vacía.

    `False` y `0` son valores llenos — tratarlos como vacíos descarta datos buenos.
    """
    return v is None or (isinstance(v, str) and v.strip() == "")


def fila_util(fila):
    """Descarta las filas de una sola celda (títulos, separadores)."""
    return sum(0 if vacia(v) else 1 for v in fila) > 1


def cargar(hoja="Colegios", solo_colegios=True):
    """Devuelve (encabezados, [dict por fila]) ya saneado.

    Con `solo_colegios` descarta además las filas sin `Colegio`. Hace falta:
    la hoja trae 885 filas de residuo con solo `Dominio web` y/o `Facebook`,
    que `fila_util` NO atrapa porque tienen más de una celda llena. Sin este
    filtro cualquier conteo se infla ~15%.
    """
    import openpyxl
    wb = openpyxl.load_workbook(ruta_base(), read_only=True, data_only=True)
    if hoja not in wb.sheetnames:
        raise SystemExit(f"No existe la hoja {hoja!r}. Hay: {wb.sheetnames}")
    ws = wb[hoja]
    h = fila_encabezado(ws)
    filas = list(ws.iter_rows(min_row=h, values_only=True))
    cols = [c for c in filas[0]]
    datos = [dict(zip(cols, f)) for f in filas[1:] if fila_util(f)]
    wb.close()
    if solo_colegios and "Colegio" in cols:
        con_nombre = [d for d in datos if not vacia(d.get("Colegio"))]
        descartadas = len(datos) - len(con_nombre)
        if descartadas:
            print(f"[aviso] {descartadas} filas sin 'Colegio' descartadas "
                  f"(residuo). Usa solo_colegios=False para incluirlas.",
                  file=sys.stderr)
        datos = con_nombre
    return cols, datos


def respaldar(fase):
    """Copia la base viva a …pre-<fase>.xlsx. Se niega a pisar un respaldo previo."""
    src = ruta_base()
    dst = src.replace(".xlsx", f".pre-{fase}.xlsx")
    if os.path.exists(dst):
        raise SystemExit(f"Ya existe {os.path.basename(dst)} — no lo piso.\n"
                         f"Usa otro nombre de fase o muévelo tú.")
    shutil.copy2(src, dst)
    return dst


def main():
    cmd = sys.argv[1] if len(sys.argv) > 1 else "info"

    if cmd == "info":
        import openpyxl
        f = ruta_base()
        print("ARCHIVO :", os.path.basename(f))
        print("RUTA    :", f)
        wb = openpyxl.load_workbook(f, read_only=True)
        for ws in wb.worksheets:
            h = fila_encabezado(ws)
            print(f"  hoja {ws.title!r}: {ws.max_row - h} filas de datos, "
                  f"{ws.max_column} columnas (encabezado en fila {h})")
        wb.close()
        print("\nRespaldos existentes:")
        # El original es .xlsx.bak, no .xlsx — hay que buscar los dos patrones.
        for r in sorted(glob.glob(os.path.join(RAIZ, PATRON)) +
                        glob.glob(os.path.join(RAIZ, PATRON + ".bak"))):
            b = os.path.basename(r)
            if any(s in b for s in RESPALDOS):
                print("  ", b)

    elif cmd == "cols":
        hoja = sys.argv[2] if len(sys.argv) > 2 else "Colegios"
        cols, datos = cargar(hoja)
        llenas = {c: sum(1 for d in datos if not vacia(d.get(c))) for c in cols}
        print(f"{hoja}: {len(datos)} filas\n")
        for c in cols:
            n = llenas[c]
            print(f"  {str(c)[:34]:<34} {n:>6}  {n / len(datos) * 100:5.1f}%")

    elif cmd == "respaldar":
        if len(sys.argv) < 3:
            raise SystemExit("Falta el nombre de la fase: base.py respaldar <fase>")
        print("Respaldo creado:", os.path.basename(respaldar(sys.argv[2])))

    elif cmd == "contar":
        hoja = sys.argv[2] if len(sys.argv) > 2 else "Colegios"
        cols, datos = cargar(hoja)
        if len(sys.argv) > 3 and "=" in sys.argv[3]:
            col, val = sys.argv[3].split("=", 1)
            n = sum(1 for d in datos if _norm(d.get(col)) == _norm(val))
            print(f"{hoja} donde {col}={val!r}: {n} de {len(datos)}")
        else:
            print(f"{hoja}: {len(datos)} filas")

    else:
        raise SystemExit(__doc__)


if __name__ == "__main__":
    main()
